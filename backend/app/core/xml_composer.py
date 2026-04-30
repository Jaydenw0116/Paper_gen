import re
from io import BytesIO
from docx import Document
from docxcompose.composer import Composer
from docx.oxml.ns import qn
from typing import List, Dict

class DocumentComposer:
    QUESTION_NUMBER_REGEX = re.compile(r'^\s*\d+[\.．、](?!\d)')
    
    @staticmethod
    def renumber_question(doc_bytes: bytes, new_number: int) -> bytes:
        doc = Document(BytesIO(doc_bytes))
        
        for element in doc.element.body:
            if element.tag.endswith('p'):
                DocumentComposer._renumber_paragraph(element, new_number)
        
        output = BytesIO()
        doc.save(output)
        return output.getvalue()

    @staticmethod
    def _renumber_paragraph(paragraph_element, new_number: int):
        for child in paragraph_element.iter():
            tag = child.tag
            if tag.endswith('drawing') or tag.endswith('shape') or 'AlternateContent' in tag:
                continue
            
            if tag.endswith('t') and child.text:
                match = DocumentComposer.QUESTION_NUMBER_REGEX.match(child.text)
                if match:
                    new_prefix = f"{new_number}."
                    child.text = new_prefix + child.text[len(match.group()):]
                    break

    @staticmethod
    def compose_documents(template_bytes: bytes, question_bytes_list: List[bytes]) -> bytes:
        template_doc = Document(BytesIO(template_bytes))
        composer = Composer(template_doc)
        
        for q_bytes in question_bytes_list:
            question_doc = Document(BytesIO(q_bytes))
            
            if not question_doc.sections:
                question_doc.add_section()
            
            composer.append(question_doc)
        
        output = BytesIO()
        composer.save(output)
        return output.getvalue()

    @staticmethod
    def create_empty_template() -> bytes:
        doc = Document()
        output = BytesIO()
        doc.save(output)
        return output.getvalue()