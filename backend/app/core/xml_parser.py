import re
import uuid
import base64
from io import BytesIO
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement, parse_xml
from docx.shared import Inches
from typing import List, Dict, Any, Tuple

class QuestionParser:
    QUESTION_NUMBER_REGEX = re.compile(r'^\s*\d+[\.．、](?!\d+$)|^\s*\d+[\.．、]$')

    IGNORED_CONTENT_REGEX = re.compile(
        r'^\s*(多项选择题|选择题|填空题|判断题|问答题|解答题|计算题|证明题|评分标准|答案|参考答案|得分|姓名|班级|学号|密封线)\s*$',
        re.IGNORECASE
    )

    TABLE_ANSWER_PATTERN = re.compile(
        r'^\s*(答案|选项|A|B|C|D|(\d+)\s*[、.．])\s*$',
        re.IGNORECASE
    )

    DOCUMENT_END_PATTERNS = re.compile(
        r'\s*(小測完|小测完|考试完|考試完|測驗完|测验完)\s*$',
        re.IGNORECASE
    )

    SECTION_HEADER_PATTERNS = re.compile(
        r'多[项項][选選][择擇][题題]|填空[题題]|判断[题題]|问答[题題]|解答[题題]|计算[题題]|证明[题題]|附加[题題]|评分标准',
        re.IGNORECASE
    )

    QUESTION_INDICATOR_PATTERNS = re.compile(
        r'(下列各[題题]|下列[题題]目|每[题題]|每小题|本题|本题占|[佔]\s*[（(]?\s*分)',
        re.IGNORECASE
    )

    ANSWER_CHOICE_PATTERN = re.compile(
        r'^[A-D][.．、、]',
        re.IGNORECASE
    )

    @staticmethod
    def parse_document(doc_bytes: bytes) -> List[Dict[str, Any]]:
        doc = Document(BytesIO(doc_bytes))
        questions = []

        current_question = []
        current_number = None
        question_start = 0

        for i, element in enumerate(doc.element.body):
            if element.tag.endswith('p'):
                text = QuestionParser._get_element_text(element).strip()

                if QuestionParser.IGNORED_CONTENT_REGEX.match(text):
                    if current_question:
                        questions.append({
                            'start': question_start,
                            'end': i,
                            'number': current_number,
                            'text': '\n'.join(current_question)
                        })
                        current_question = []
                        current_number = None
                    continue

                # 检查段落内部是否包含多个题号（处理表格内容合并的情况）
                parts = QuestionParser._split_text_by_question_numbers(text)
                
                for j, part in enumerate(parts):
                    part_text = part['text']
                    part_has_number = part['has_number']
                    
                    if part_has_number:
                        # 这部分以题号开头
                        if current_question:
                            questions.append({
                                'start': question_start,
                                'end': i,
                                'number': current_number,
                                'text': '\n'.join(current_question)
                            })
                        current_question = [part_text]
                        current_number = part['number']
                        question_start = i
                    elif current_question:
                        current_question.append(part_text)
                    # 如果没有题号且没有当前问题，跳过
            elif element.tag.endswith('tbl') and current_question:
                if not QuestionParser._is_answer_table(element):
                    current_question.append('[TABLE]')
            elif element.tag.endswith('tbl') and current_question:
                if not QuestionParser._is_answer_table(element):
                    current_question.append('[TABLE]')

        if current_question:
            questions.append({
                'start': question_start,
                'end': len(doc.element.body),
                'number': current_number,
                'text': '\n'.join(current_question)
            })

        questions = QuestionParser._validate_questions(questions)
        return questions

    @staticmethod
    def _validate_questions(questions: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
        validated = []
        for q in questions:
            text = q.get('text', '')

            if not text or len(text.strip()) < 5:
                continue

            lines = text.split('\n')
            if not lines:
                continue

            first_line = lines[0].strip()
            number_match = QuestionParser.QUESTION_NUMBER_REGEX.match(first_line)
            if number_match:
                first_line_content = first_line[len(number_match.group()):].strip()
                if len(first_line_content) < 3 and len(lines) <= 2:
                    # 允许只有题号的问题（如选择题答案）
                    if len(lines) == 1 and first_line_content == '':
                        # 单一行且只有题号，保留
                        pass
                    else:
                        continue

            question_text = text

            lines = question_text.split('\n')
            valid_lines = []
            for line in lines:
                if QuestionParser.DOCUMENT_END_PATTERNS.search(line):
                    break
                valid_lines.append(line)

            if valid_lines:
                question_text = '\n'.join(valid_lines)

            if len(question_text.strip()) < 5:
                continue

            question_text = QuestionParser._clean_question_text(question_text)

            if len(question_text.strip()) < 5:
                continue

            if 'TABLE' in question_text:
                non_table_lines = [l for l in question_text.split('\n') if '[TABLE]' not in l]
                text_only = ''.join(non_table_lines)
                text_only = re.sub(r'\s+', '', text_only)
                if len(text_only) < 5:
                    continue

            q['text'] = question_text
            validated.append(q)

        return validated

    @staticmethod
    def _clean_question_text(text: str) -> str:
        lines = text.split('\n')
        valid_lines = []
        skip_mode = False

        for line in lines:
            stripped = line.strip()

            if QuestionParser.SECTION_HEADER_PATTERNS.search(stripped):
                skip_mode = True
                continue

            if skip_mode:
                if QuestionParser.QUESTION_INDICATOR_PATTERNS.search(stripped):
                    continue
                if QuestionParser.ANSWER_CHOICE_PATTERN.match(stripped):
                    skip_mode = False

            if stripped:
                valid_lines.append(line)

        return '\n'.join(valid_lines)

    @staticmethod
    def _is_answer_table(element) -> bool:
        text_content = []
        for tbl in element.iter():
            if tbl.tag.endswith('t') and tbl.text:
                text_content.append(tbl.text.strip())

        full_text = ' '.join(text_content)
        if QuestionParser.TABLE_ANSWER_PATTERN.search(full_text):
            return True
        if len(text_content) > 0 and all(len(t) <= 5 for t in text_content):
            return True
        return False

    @staticmethod
    def _get_element_text(element) -> str:
        text_parts = []
        for child in element.iter():
            if child.tag.endswith('t'):
                text_parts.append(child.text or '')
            elif child.tag.endswith('tab'):
                text_parts.append(' ')
            elif child.tag.endswith('br'):
                text_parts.append('\n')
        return ''.join(text_parts)

    @staticmethod
    def _split_text_by_question_numbers(text: str) -> list:
        """
        将文本按题号分割，处理段落内部包含多个题号的情况
        返回格式: [{'text': 'xxx', 'has_number': True/False, 'number': '1.'}]
        """
        if not text:
            return [{'text': '', 'has_number': False, 'number': ''}]
        
        parts = []
        
        # 查找文本中所有题号位置
        number_pattern = re.compile(r'(?<!\d)\d+[\.．、](?!\d)')
        matches = list(number_pattern.finditer(text))
        
        if not matches:
            # 没有找到任何题号
            return [{'text': text, 'has_number': False, 'number': ''}]
        
        # 如果第一个题号不在开头，先添加开头的内容
        if matches[0].start() > 0:
            parts.append({'text': text[:matches[0].start()].strip(), 'has_number': False, 'number': ''})
        
        # 处理每个题号
        for i, match in enumerate(matches):
            # 获取当前题号
            current_number = match.group().strip()
            
            # 确定下一个题号的位置
            if i < len(matches) - 1:
                next_match = matches[i + 1]
                end_pos = next_match.start()
            else:
                end_pos = len(text)
            
            # 获取从当前题号到下一个题号之间的内容
            content = text[match.start():end_pos].strip()
            parts.append({'text': content, 'has_number': True, 'number': current_number})
        
        # 清理空的部分
        return [p for p in parts if p['text'].strip()]


    @staticmethod
    def extract_question_block(doc_bytes: bytes, start_idx: int, end_idx: int, question_number: str = None) -> bytes:
        doc = Document(BytesIO(doc_bytes))
        body_len = len(doc.element.body)

        end_idx = min(end_idx + 1, body_len)
        start_idx = max(start_idx, 0)

        nodes_to_remove = []

        for i in range(body_len - 1, end_idx - 1, -1):
            if i >= 0 and i < len(doc.element.body):
                if not doc.element.body[i].tag.endswith('sectPr'):
                    nodes_to_remove.append(doc.element.body[i])

        for i in range(start_idx - 1, -1, -1):
            if i >= 0 and i < len(doc.element.body):
                if not doc.element.body[i].tag.endswith('sectPr'):
                    nodes_to_remove.append(doc.element.body[i])

        for node in nodes_to_remove:
            try:
                if node in doc.element.body:
                    doc.element.body.remove(node)
            except Exception:
                pass

        if len(doc.element.body) > 0:
            last_child = doc.element.body[-1]
            if last_child is not None and not last_child.tag.endswith('sectPr'):
                new_p = OxmlElement('w:p')
                doc.element.body.insert(len(doc.element.body) - 1, new_p)

        output = BytesIO()
        doc.save(output)
        return output.getvalue()


    @staticmethod
    def extract_preview_text(doc_bytes: bytes) -> str:
        doc = Document(BytesIO(doc_bytes))
        text = []

        for element in doc.element.body:
            if element.tag.endswith('p'):
                text.append(QuestionParser._get_element_text(element))
            elif element.tag.endswith('tbl'):
                text.append('[TABLE]')

        full_text = '\n'.join(text)
        full_text = re.sub(r'\b\d{9,}\b', '', full_text)
        return full_text.strip()[:500]

    @staticmethod
    def extract_images(doc_bytes: bytes) -> List[str]:
        images = []
        doc = Document(BytesIO(doc_bytes))

        for rel in doc.part.rels.values():
            if 'image' in rel.target_ref:
                image_data = rel.target_part.blob
                if image_data:
                    base64_img = base64.b64encode(image_data).decode('utf-8')
                    images.append(base64_img)

        return images
