from fastapi import APIRouter, File, UploadFile, HTTPException, BackgroundTasks
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from typing import List
import uuid
from io import BytesIO
from docx import Document

from ..core.xml_parser import QuestionParser
from ..core.xml_composer import DocumentComposer
from ..core.cache import cache_manager
from ..core.config import settings

router = APIRouter()

@router.post("/clear-cache")
async def clear_cache():
    """清空所有缓存数据"""
    if hasattr(cache_manager.client, '_data'):
        cache_manager.client._data.clear()
    return {"message": "缓存已清空"}

class QuestionItem(BaseModel):
    id: str
    original_no: str
    preview_text: str
    images: List[str]
    source_file: str = ""

class ComposeRequest(BaseModel):
    question_ids: List[str]
    title: str = "试卷"
    subject: str = ""
    total_score: int = 100

class UploadToBankRequest(BaseModel):
    session_id: str
    question_ids: List[str]

@router.post("/upload", response_model=dict)
async def upload_documents(
    question_file: UploadFile = File(...),
    answer_file: UploadFile = File(...)
):
    if question_file.content_type not in ["application/vnd.openxmlformats-officedocument.wordprocessingml.document", "application/octet-stream"]:
        raise HTTPException(status_code=400, detail="题目文档必须是.docx格式")
    
    if answer_file.content_type not in ["application/vnd.openxmlformats-officedocument.wordprocessingml.document", "application/octet-stream"]:
        raise HTTPException(status_code=400, detail="答案文档必须是.docx格式")
    
    question_content = await question_file.read()
    answer_content = await answer_file.read()
    
    if len(question_content) > settings.max_file_size:
        raise HTTPException(status_code=400, detail="题目文档大小超过限制")
    
    if len(answer_content) > settings.max_file_size:
        raise HTTPException(status_code=400, detail="答案文档大小超过限制")
    
    questions = QuestionParser.parse_document(question_content)
    answers = QuestionParser.parse_document(answer_content)
    
    question_items = []
    session_id = cache_manager.generate_session_id()
    source_file = question_file.filename
    
    for idx, q in enumerate(questions):
        q_id = str(uuid.uuid4())
        
        # 使用问题文本直接创建文档，避免提取块时包含其他题号
        if q.get('text'):
            q_doc = Document()
            q_doc.add_paragraph(q['text'])
            q_output = BytesIO()
            q_doc.save(q_output)
            q_bytes = q_output.getvalue()
        else:
            q_bytes = QuestionParser.extract_question_block(question_content, q['start'], q['end'])
        
        preview_text = q.get('text', '')[:500] if q.get('text') else QuestionParser.extract_preview_text(q_bytes)
        images = QuestionParser.extract_images(q_bytes)
        
        cache_manager.set_docx_bytes(f"session:{session_id}:question:{q_id}", q_bytes)
        
        if idx < len(answers):
            ans = answers[idx]
            if ans.get('text'):
                ans_doc = Document()
                ans_doc.add_paragraph(ans['text'])
                ans_output = BytesIO()
                ans_doc.save(ans_output)
                answer_bytes = ans_output.getvalue()
            else:
                answer_bytes = QuestionParser.extract_question_block(answer_content, ans['start'], ans['end'])
        else:
            answer_bytes = q_bytes
        
        cache_manager.set_docx_bytes(f"session:{session_id}:answer:{q_id}", answer_bytes)
        
        question_items.append({
            "id": q_id,
            "original_no": q['number'],
            "preview_text": preview_text,
            "images": images,
            "source_file": source_file
        })
    
    cache_manager.set_questions(session_id, question_items)
    cache_manager.expire_session(session_id, 3600)
    
    return {"session_id": session_id, "questions": question_items, "source_file": source_file}

@router.post("/upload-to-bank")
async def upload_to_bank(request: UploadToBankRequest):
    bank_data = cache_manager.get_docx_bytes("question_bank")
    if bank_data is None:
        bank = {"questions": [], "answers": {}}
    else:
        import ast
        if isinstance(bank_data, bytes):
            bank_str = bank_data.decode('utf-8')
        else:
            bank_str = str(bank_data)
        bank = ast.literal_eval(bank_str)
    
    questions = cache_manager.get_questions(request.session_id)
    if not questions:
        raise HTTPException(status_code=404, detail="会话不存在或已过期")
    
    added_count = 0
    for q_id in request.question_ids:
        q_bytes = cache_manager.get_docx_bytes(f"session:{request.session_id}:question:{q_id}")
        a_bytes = cache_manager.get_docx_bytes(f"session:{request.session_id}:answer:{q_id}")
        
        if q_bytes and a_bytes:
            q_info = next((q for q in questions if q['id'] == q_id), None)
            bank["questions"].append({
                "id": q_id,
                "source_file": q_info.get("source_file", ""),
                "original_no": q_info.get("original_no", "")
            })
            bank["answers"][q_id] = {"question": q_bytes, "answer": a_bytes}
            added_count += 1
    
    cache_manager.set_docx_bytes("question_bank", str(bank))
    return {"message": "题目已添加到题库", "count": added_count}

@router.get("/bank", response_model=List[QuestionItem])
async def get_question_bank():
    bank_data = cache_manager.get_docx_bytes("question_bank")
    if bank_data is None:
        return []
    
    import ast
    if isinstance(bank_data, bytes):
        bank_str = bank_data.decode('utf-8')
    else:
        bank_str = str(bank_data)
    bank = ast.literal_eval(bank_str)
    
    questions = []
    for q in bank.get("questions", []):
        q_data = bank["answers"].get(q["id"], {})
        q_bytes = q_data.get("question")
        if q_bytes:
            preview_text = QuestionParser.extract_preview_text(q_bytes)
            images = QuestionParser.extract_images(q_bytes)
            questions.append({
                "id": q["id"],
                "original_no": q.get("original_no", ""),
                "preview_text": preview_text,
                "images": images,
                "source_file": q.get("source_file", "")
            })
    
    return questions

@router.delete("/bank/{question_id}")
async def remove_from_bank(question_id: str):
    bank_data = cache_manager.get_docx_bytes("question_bank")
    if bank_data is None:
        raise HTTPException(status_code=404, detail="题库为空")
    
    import ast
    if isinstance(bank_data, bytes):
        bank_str = bank_data.decode('utf-8')
    else:
        bank_str = str(bank_data)
    bank = ast.literal_eval(bank_str)
    
    bank["questions"] = [q for q in bank.get("questions", []) if q["id"] != question_id]
    bank["answers"].pop(question_id, None)
    
    cache_manager.set_docx_bytes("question_bank", str(bank))
    return {"message": "题目已从题库删除"}

@router.delete("/bank")
async def clear_bank():
    cache_manager.set_docx_bytes("question_bank", str({"questions": [], "answers": {}}))
    return {"message": "题库已清空"}

@router.get("/questions/{session_id}", response_model=List[QuestionItem])
async def get_questions(session_id: str):
    questions = cache_manager.get_questions(session_id)
    if not questions:
        raise HTTPException(status_code=404, detail="会话不存在或已过期")
    return questions

@router.post("/compose")
async def compose_documents(request: ComposeRequest):
    bank_data = cache_manager.get_docx_bytes("question_bank")
    if bank_data is None:
        raise HTTPException(status_code=404, detail="题库为空")

    import ast
    if isinstance(bank_data, bytes):
        bank_str = bank_data.decode('utf-8')
    else:
        bank_str = str(bank_data)
    bank = ast.literal_eval(bank_str)

    if len(request.question_ids) > settings.max_questions:
        raise HTTPException(status_code=400, detail="单次组卷题目数量超过限制")

    question_bytes_list = []
    answer_bytes_list = []

    for q_id in request.question_ids:
        q_data = bank["answers"].get(q_id)
        if not q_data:
            raise HTTPException(status_code=404, detail=f"题目ID {q_id} 不存在于题库中")

        question_bytes_list.append(q_data["question"])
        answer_bytes_list.append(q_data["answer"])

    import os
    project_root = os.path.dirname(os.path.dirname(os.path.dirname(os.path.dirname(__file__))))
    template_path = os.path.join(project_root, "template.docx")
    template_ans_path = os.path.join(project_root, "template_ans.docx")

    try:
        paper_template = DocumentComposer.load_template(template_path)
    except Exception:
        paper_template = DocumentComposer.create_empty_template()

    try:
        answer_template = DocumentComposer.load_template(template_ans_path)
    except Exception:
        answer_template = DocumentComposer.create_empty_template()

    paper_doc = DocumentComposer.compose_documents(paper_template, question_bytes_list)
    answer_doc = DocumentComposer.compose_documents(answer_template, answer_bytes_list)

    session_id = cache_manager.generate_session_id()
    cache_manager.set_docx_bytes(f"session:{session_id}:paper", paper_doc)
    cache_manager.set_docx_bytes(f"session:{session_id}:answer", answer_doc)

    return {"message": "组卷成功", "session_id": session_id}

@router.get("/download/{session_id}/{doc_type}")
async def download_document(session_id: str, doc_type: str):
    if doc_type not in ["paper", "answer"]:
        raise HTTPException(status_code=400, detail="无效的文档类型")
    
    doc_bytes = cache_manager.get_docx_bytes(f"session:{session_id}:{doc_type}")
    if not doc_bytes:
        raise HTTPException(status_code=404, detail="文档不存在或已过期")
    
    filename = "试卷.docx" if doc_type == "paper" else "答案.docx"
    import urllib.parse
    encoded_filename = urllib.parse.quote(filename)
    
    return StreamingResponse(
        BytesIO(doc_bytes),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{encoded_filename}"}
    )

@router.delete("/session/{session_id}")
async def delete_session(session_id: str):
    cache_manager.delete_session(session_id)
    return {"message": "会话已删除"}

@router.post("/debug-answer", response_model=dict)
async def debug_answer(
    answer_file: UploadFile = File(...)
):
    answer_content = await answer_file.read()
    answers = QuestionParser.parse_document(answer_content)
    
    result = []
    for ans in answers:
        result.append({
            "number": ans.get('number'),
            "start": ans.get('start'),
            "end": ans.get('end'),
            "preview": ans.get('text', '')[:50]
        })
    
    return {"answers": result}
