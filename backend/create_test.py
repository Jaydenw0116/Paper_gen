from docx import Document

doc = Document()
doc.add_paragraph("1. 第一道题目")
doc.add_paragraph("这是第一道题的内容，包含一些文字说明。")
doc.add_paragraph("2. 第二道题目")
doc.add_paragraph("这是第二道题的内容，用于测试切割功能。")
doc.add_paragraph("3. 第三道题目")
doc.add_paragraph("这是第三道题的内容，测试完整的流程。")
doc.save('test_question.docx')

doc2 = Document()
doc2.add_paragraph("1. 第一题答案")
doc2.add_paragraph("这是第一题的答案。")
doc2.add_paragraph("2. 第二题答案")
doc2.add_paragraph("这是第二题的答案。")
doc2.add_paragraph("3. 第三题答案")
doc2.add_paragraph("这是第三题的答案。")
doc2.save('test_answer.docx')

print("测试文档创建成功！")
