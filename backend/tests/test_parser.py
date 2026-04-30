import pytest
import os
from io import BytesIO
from docx import Document
from app.core.xml_parser import QuestionParser
from app.core.xml_composer import DocumentComposer

TEST_DATA_DIR = os.path.join(os.path.dirname(__file__), 'test_data')

def create_test_doc(lines):
    doc = Document()
    for line in lines:
        if line.strip():
            doc.add_paragraph(line)
        else:
            doc.add_paragraph()
    output = BytesIO()
    doc.save(output)
    return output.getvalue()

def test_parse_document():
    content = [
        "1. 第一题",
        "这是第一题的内容",
        "",
        "2. 第二题",
        "这是第二题的内容",
        "",
        "3. 第三题",
        "这是第三题的内容"
    ]
    
    doc_bytes = create_test_doc(content)
    questions = QuestionParser.parse_document(doc_bytes)
    
    assert len(questions) == 3
    assert questions[0]['number'] == '1.'
    assert questions[1]['number'] == '2.'
    assert questions[2]['number'] == '3.'

def test_extract_preview_text():
    content = ["1. 测试题目这是一段很长的文本内容，用于测试预览功能，确保能够正确截取前500个字符。"]
    doc_bytes = create_test_doc(content)
    preview = QuestionParser.extract_preview_text(doc_bytes)
    
    assert '1. 测试题目' in preview
    assert len(preview) <= 500

def test_renumber_question():
    content = ["1. 原始题目内容"]
    doc_bytes = create_test_doc(content)
    
    renumbered = DocumentComposer.renumber_question(doc_bytes, 5)
    doc = Document(BytesIO(renumbered))
    text = doc.paragraphs[0].text
    
    assert text.startswith('5.')
    assert '原始题目内容' in text

def test_compose_documents():
    template = DocumentComposer.create_empty_template()
    question1 = create_test_doc(["1. 题目一"])
    question2 = create_test_doc(["2. 题目二"])
    
    composed = DocumentComposer.compose_documents(template, [question1, question2])
    doc = Document(BytesIO(composed))
    
    assert len(doc.paragraphs) >= 2
    assert '题目一' in doc.paragraphs[0].text or '题目一' in doc.paragraphs[1].text

def test_regex_match():
    test_cases = [
        ('1. 题目', True),
        ('  2．题目', True),
        ('3、题目', True),
        ('123.45', False),
        ('abc1.题目', False),
    ]
    
    for test, expected in test_cases:
        match = QuestionParser.QUESTION_NUMBER_REGEX.match(test)
        assert (match is not None) == expected

if __name__ == '__main__':
    pytest.main([__file__, '-v'])